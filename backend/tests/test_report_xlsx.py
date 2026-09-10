"""Content of the hours spreadsheet: editable values and formula totals."""
import io
import uuid
from datetime import date, timedelta
from decimal import Decimal

from openpyxl import load_workbook

from app.schemas.informes import HorasEntryRow, HorasReportOut, HorasRow, TradeHoursRow
from app.services.report_xlsx import build_horas_xlsx

OBRA_A, OBRA_B = uuid.uuid4(), uuid.uuid4()
ANA, LUIS = uuid.uuid4(), uuid.uuid4()


def entry(user_id, name, work_date, hours, *, notes=None, obra_id=OBRA_A,
          obra="Obra Centro", trade="Fontanero", validated=True, edited=False):
    return HorasEntryRow(
        id=uuid.uuid4(), obra_id=obra_id, obra_name=obra, user_id=user_id,
        user_full_name=name, trade=trade, work_date=work_date,
        hours=Decimal(hours), notes=notes, validated=validated,
        edited_by_admin=edited, media_count=0,
    )


def report(entries) -> HorasReportOut:
    rows: dict[tuple, HorasRow] = {}
    for e in entries:
        key = (e.obra_id, e.user_id)
        if key not in rows:
            rows[key] = HorasRow(
                obra_id=e.obra_id, obra_name=e.obra_name, user_id=e.user_id,
                user_full_name=e.user_full_name, trade=e.trade,
                total_hours=Decimal("0"), entry_count=0,
            )
        rows[key].total_hours += e.hours
        rows[key].entry_count += 1
    trades = sorted({e.trade for e in entries}, key=lambda t: (t is None, t or ""))
    return HorasReportOut(
        rows=list(rows.values()), entries=entries,
        by_trade=[TradeHoursRow(trade=t, total_hours=Decimal("0")) for t in trades],
        total_hours=sum((e.hours for e in entries), Decimal("0")),
        total_entries=len(entries),
    )


def build(entries, *, data_only=False):
    """data_only=True reads the stored formula results, as a viewer that doesn't recalculate would."""
    content = build_horas_xlsx(
        report(entries), period="01/08/2026 – 31/08/2026",
        obra_label=None, worker_label=None, generated="10/09/2026 09:00",
    )
    return load_workbook(io.BytesIO(content), data_only=data_only)


def values(ws) -> list[list]:
    return [list(r) for r in ws.iter_rows(values_only=True)]


def find_row(ws, first_cell) -> int:
    for row in ws.iter_rows(min_col=1, max_col=1):
        if row[0].value == first_cell:
            return row[0].row
    raise AssertionError(f"{first_cell!r} not found")


def test_sheets():
    wb = build([entry(ANA, "Ana", date(2026, 8, 3), "8")])
    assert wb.sheetnames == ["Resumen", "Detalle"]


def test_detail_has_one_editable_row_per_parte():
    wb = build([
        entry(LUIS, "Luis", date(2026, 8, 5), "6", trade=None),
        entry(ANA, "Ana", date(2026, 8, 4), "7.5", notes="Cambio de bajante"),
        entry(ANA, "Ana", date(2026, 8, 3), "8.63", validated=False, edited=True),
    ])
    ws = wb["Detalle"]
    rows = values(ws)

    assert rows[0] == ["Fecha", "Obra", "Trabajador", "Oficio", "Horas",
                       "Validado", "Editado por admin", "Notas"]
    # Grouped by worker, then chronological
    assert [(r[2], r[0].date()) for r in rows[1:]] == [
        ("Ana", date(2026, 8, 3)), ("Ana", date(2026, 8, 4)), ("Luis", date(2026, 8, 5)),
    ]
    # Hours are real durations (read back as timedelta), rounded to the minute
    assert rows[1][4] == timedelta(minutes=518)
    assert rows[2][4] == timedelta(hours=7, minutes=30)
    assert ws["E2"].number_format == "[h]:mm"
    assert ws["A2"].number_format == "dd/mm/yyyy"
    assert rows[1][5:7] == ["Pendiente", "Sí"]
    assert rows[2][7] == "Cambio de bajante"
    assert rows[3][3] == "Sin oficio"
    assert ws.freeze_panes == "A2"


def test_summary_totals_are_formulas_over_the_detail():
    wb = build([
        entry(ANA, "Ana", date(2026, 8, 3), "8"),
        entry(ANA, "Ana", date(2026, 8, 4), "4", obra_id=OBRA_B, obra="Nave"),
        entry(LUIS, "Luis", date(2026, 8, 5), "6", trade="Electricista"),
    ])
    ws = wb["Resumen"]

    assert ws["B5"].value == "=SUM(Detalle!$E:$E)"
    assert ws["B6"].value == "=COUNT(Detalle!$A:$A)"

    r = find_row(ws, "Obra Centro")
    assert ws[f"B{r}"].value == "Ana"
    assert ws[f"D{r}"].value == f"=COUNTIFS(Detalle!$B:$B,A{r},Detalle!$C:$C,B{r})"
    assert ws[f"E{r}"].value == f"=SUMIFS(Detalle!$E:$E,Detalle!$B:$B,A{r},Detalle!$C:$C,B{r})"

    trade = find_row(ws, "Electricista")
    assert ws[f"B{trade}"].value == f"=SUMIFS(Detalle!$E:$E,Detalle!$D:$D,A{trade})"

    header = find_row(ws, "Trabajador")
    assert [ws[f"A{header + 1}"].value, ws[f"A{header + 2}"].value] == ["Ana", "Luis"]
    total = header + 3
    assert ws[f"A{total}"].value == "TOTAL"
    assert ws[f"D{total}"].value == f"=SUM(D{header + 1}:D{header + 2})"


def test_formulas_store_their_result_for_viewers_that_do_not_recalculate():
    wb = build([
        entry(ANA, "Ana", date(2026, 8, 3), "8"),
        entry(ANA, "Ana", date(2026, 8, 4), "4", obra_id=OBRA_B, obra="Nave"),
        entry(LUIS, "Luis", date(2026, 8, 5), "6.5", trade="Electricista"),
    ], data_only=True)
    ws = wb["Resumen"]

    assert ws["B5"].value == timedelta(hours=18, minutes=30)
    assert ws["B6"].value == 3

    r = find_row(ws, "Obra Centro")
    assert ws[f"D{r}"].value == 1
    assert ws[f"E{r}"].value == timedelta(hours=8)
    assert ws[f"B{find_row(ws, 'Electricista')}"].value == timedelta(hours=6, minutes=30)
    assert ws[f"B{find_row(ws, 'Fontanero')}"].value == timedelta(hours=12)

    header = find_row(ws, "Trabajador")
    assert ws[f"D{header + 1}"].value == timedelta(hours=12)  # Ana, both obras
    totals = [c.row for c in ws["A"] if c.value == "TOTAL"]
    assert len(totals) == 2
    for t in totals:
        assert timedelta(hours=18, minutes=30) in (ws[f"D{t}"].value, ws[f"E{t}"].value)


def test_text_that_looks_like_a_formula_stays_text():
    wb = build([
        entry(ANA, "Ana", date(2026, 8, 3), "8", notes='=HYPERLINK("http://x","y")',
              obra="=1+1"),
    ])
    ws = wb["Detalle"]
    assert ws["H2"].data_type == "s"
    assert ws["H2"].value == '=HYPERLINK("http://x","y")'
    assert ws["B2"].data_type == "s"


def test_control_characters_are_dropped():
    wb = build([entry(ANA, "Ana", date(2026, 8, 3), "8", notes="Fuga\x0b grave")])
    assert wb["Detalle"]["H2"].value == "Fuga grave"


def test_empty_report():
    wb = build([])
    assert values(wb["Detalle"]) == [["Fecha", "Obra", "Trabajador", "Oficio", "Horas",
                                      "Validado", "Editado por admin", "Notas"]]
    assert find_row(wb["Resumen"], "Sin datos para estos filtros")
